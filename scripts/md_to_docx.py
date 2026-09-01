#!/usr/bin/env python3
"""Convert a generated proposal Markdown into a deliverable Word (.docx).

Uses pandoc (bundled via pypandoc_binary) for high-fidelity conversion:
- headings / pipe tables / images
- LaTeX math (\\( \\) inline, \\[ \\] display) -> native Word equations (OMML)
- a Chinese reference template (宋体 body, 黑体 headings) for styling

Usage:
  python scripts/md_to_docx.py outputs/task_001_v4/proposal_final.md
  python scripts/md_to_docx.py <md> --out foo.docx --ref assets/reference.docx
Images referenced as figures/xxx.png are resolved relative to the md's folder.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

_ALIGN = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
          "center": WD_ALIGN_PARAGRAPH.CENTER}

# ── Formatting for a Chinese proposal (tweak here) ──────────────────────────
#   size: pt | indent_chars: 首行缩进(字符) | line: 行距倍数
#   before/after: 段前/段后(pt) | align: justify/left/center
_BODY = dict(latin="Times New Roman", ea="宋体", size=12, bold=False,
             indent_chars=2, line=1.5, before=0, after=0, align="justify")
_H = lambda size, before, after, align="left": dict(
    latin="Times New Roman", ea="黑体", size=size, bold=True,
    indent_chars=0, line=1.5, before=before, after=after, align=align)

_STYLES = {
    "Normal":          _BODY,
    "Body Text":       _BODY,
    "First Paragraph": _BODY,
    "Compact":         dict(_BODY, after=0),
    "List Paragraph":  _BODY,   # numbered/bulleted items share the body's 2-char indent
    "Title":           _H(18, 6, 14, "center"),
    "Heading 1":       _H(15, 12, 6),
    "Heading 2":       _H(14, 10, 4),
    "Heading 3":       _H(13, 8, 4),
    "Heading 4":       _H(12, 6, 2),
}


def _apply_style(style, s: dict):
    # font (Latin + East-Asian)
    style.font.size = Pt(s["size"])
    style.font.bold = s["bold"]
    style.font.italic = False   # Word's built-in Heading 4 is bold *italic* — clear it
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    # clear THEME fonts first — otherwise headings keep the template's theme font
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    rfonts.set(qn("w:ascii"), s["latin"])
    rfonts.set(qn("w:hAnsi"), s["latin"])
    rfonts.set(qn("w:eastAsia"), s["ea"])
    rfonts.set(qn("w:cs"), s["latin"])

    # force BLACK color (pandoc's default headings are themed blue); drop theme color
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    for a in ("w:themeColor", "w:themeTint", "w:themeShade"):
        if color.get(qn(a)) is not None:
            del color.attrib[qn(a)]

    # paragraph format: line spacing, spacing, alignment
    pf = style.paragraph_format
    pf.line_spacing = s["line"]
    pf.space_before = Pt(s["before"])
    pf.space_after = Pt(s["after"])
    pf.alignment = _ALIGN[s["align"]]

    # first-line indent measured in characters (真·首行缩进 N 字符)
    pPr = style.element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    chars = s["indent_chars"]
    if chars > 0:
        ind.set(qn("w:firstLineChars"), str(int(chars * 100)))        # Word 认字符
        ind.set(qn("w:firstLine"), str(int(s["size"] * chars * 20)))  # twips 兜底
    else:
        # explicit 0 — otherwise headings INHERIT the body's 2-char first-line indent
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), "0")


def _force_default_fonts(doc):
    """Make the document-wide fallback English=Times New Roman, CJK=宋体, and drop any
    monospace/theme font hardcoded on a style (e.g. pandoc's Verbatim Char=Consolas,
    or the docDefault asciiTheme=minorHAnsi→Calibri) so NO English run escapes to a
    non-Times font — even runs whose style we don't explicitly restyle."""
    # 1) doc-wide default run font
    docdef = doc.styles.element.find(qn("w:docDefaults"))
    if docdef is not None:
        rpd = docdef.find(qn("w:rPrDefault"))
        if rpd is None:
            rpd = OxmlElement("w:rPrDefault"); docdef.insert(0, rpd)
        rpr = rpd.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); rpd.append(rpr)
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        rf.set(qn("w:ascii"), "Times New Roman")
        rf.set(qn("w:hAnsi"), "Times New Roman")
        rf.set(qn("w:cs"), "Times New Roman")
        rf.set(qn("w:eastAsia"), "宋体")
    # 2) any style hardcoding a non-Times Latin font (Consolas, Calibri via theme…) -> TNR
    for s in doc.styles:
        rpr = s.element.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if rf is None:
            continue
        asc = rf.get(qn("w:ascii"))
        if (asc and asc not in ("Times New Roman", "黑体")) or rf.get(qn("w:asciiTheme")):
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            rf.set(qn("w:ascii"), "Times New Roman")
            rf.set(qn("w:hAnsi"), "Times New Roman")
            rf.set(qn("w:cs"), "Times New Roman")


def build_reference_docx(ref_path: Path) -> Path:
    """Start from pandoc's own default reference.docx (so style names match the
    writer), then restyle body/headings for a Chinese proposal."""
    ref_path.parent.mkdir(parents=True, exist_ok=True)
    pandoc = pypandoc.get_pandoc_path()
    default_bytes = subprocess.run(
        [pandoc, "--print-default-data-file", "reference.docx"],
        capture_output=True, check=True,
    ).stdout
    ref_path.write_bytes(default_bytes)

    doc = Document(str(ref_path))
    existing = {s.name: s for s in doc.styles}
    for name, spec in _STYLES.items():
        if name in existing:
            _apply_style(existing[name], spec)
    _force_default_fonts(doc)
    doc.save(str(ref_path))
    return ref_path


def _is_row(s: str) -> bool:
    return s.lstrip().startswith("|")


# a markdown list item: ordered "1." / "1)" or bullet "-" / "*" / "+"
_LIST_RE = re.compile(r"^\s*(?:\d+[.)]|[-*+])\s+")


def _is_list_item(s: str) -> bool:
    return bool(_LIST_RE.match(s))


_CN_NUM = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
           "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八", "十九", "二十"]

# an inline enumeration marker: （1）/(1) with a number inside, or ①..⑳
# (formulas use \( \) / \[ \] with letters, so bare (digit) never matches math)
_ENUM_RE = re.compile(r"(\S)[ \t]*([（(]\s*\d{1,2}\s*[)）]|[①-⑳])")

# a subheading that already carries its own ordinal — leave it alone:
#   leading "3.1 " / "一、" / "（1）"  OR  counter words 第一/目标一/方案一/成果一…
_HAS_ORDINAL = re.compile(
    r"^\s*[（(]?\s*[\d一二三四五六七八九十]+[)）]?\s*[、.．\s]"
    r"|(?:第|目标|方案|成果|任务|阶段|部分|章|节|步|批|期)[一二三四五六七八九十\d]"
)


def _break_enums(line: str) -> str:
    """Split run-together sub-points so each starts its own line/paragraph.
    Skip markers glued to a `*` — that is a bold label like **（1）…** and splitting
    it would orphan the ** (leaving a stray ** in the output)."""
    def _sub(m):
        if m.group(1) == "*":
            return m.group(0)
        return m.group(1) + "\n\n" + m.group(2)
    return _ENUM_RE.sub(_sub, line)


# a ### subheading the writer numbered with a bare Chinese numeral "一、"/"（一）"：
# this collides with the SECTION level's 一、二、, so we strip it and renumber as （一）.
_LEADING_CN_ORD = re.compile(r"^\s*[（(]?\s*([一二三四五六七八九十]{1,3})\s*[)）]?\s*[、.．]\s*")


def _number_sub(title: str, counter: list) -> str:
    """Number a ### subheading as （一）（二）…. A writer-supplied leading Chinese numeral
    ("一、发明专利…") is stripped and renumbered as （一）, so it doesn't collide with the
    section level's 一、二、; other existing ordinals (3.1 / 第一阶段 / 具体目标一) stay."""
    m = _LEADING_CN_ORD.match(title)
    if m:
        title = title[m.end():]                 # drop the writer's own "一、"
    elif _HAS_ORDINAL.search(title):
        return f"### {title}"                    # 3.1 / 第一阶段 / 具体目标一 → keep as-is
    counter[0] += 1
    n = counter[0]
    cn = _CN_NUM[n - 1] if n - 1 < len(_CN_NUM) else str(n)
    return f"### （{cn}）{title}"


# a pure leading ordinal at #### level: "1.1"/"3.4"/"1."/"1)"/"（1）"/"（一）"/"一、"
_LEADING_PURE_ORD = re.compile(
    r"^\s*(?:"
    r"\d+(?:\.\d+)+\s*"                               # 1.1 / 3.4 / 1.2.3
    r"|\d+\s*[.)、．]\s*"                             # 1. / 1) / 1、
    r"|[（(]\s*[\d一二三四五六七八九十]+\s*[)）]\s*"   # （1）/（一）
    r"|[一二三四五六七八九十]{1,3}\s*[、.．]\s*"       # 一、
    r")"
)


def _number_subsub(title: str, counter: list) -> str:
    """Number a #### sub-subheading as 1. 2. 3. within its ### parent (一、/（一）/1.
    hierarchy). A writer's pure leading ordinal (1.1 / 3.1 / （一） / 一、) is stripped and
    renumbered so it doesn't echo the parent's （一）（二）; a word-ordinal label
    (成果一 / 第一阶段) keeps its own numbering."""
    m = _LEADING_PURE_ORD.match(title)
    if m:
        title = title[m.end():]
    elif _HAS_ORDINAL.search(title):
        return f"#### {title}"
    counter[0] += 1
    return f"#### {counter[0]}. {title}"


# ── Generator artefacts: AI disclaimers and placeholder lines ───────────────
# Ported from the collaborator's document_exporter.py (see legacy/collab notes).
# Only the artefact patterns come across. That exporter also stripped Markdown
# decoration — lone asterisks, bold markers, runs of spaces — which it could
# afford because it wrote runs through python-docx and had no Markdown left to
# protect. Here it would destroy the "*图1：…*" captions and the "**[图N：…]**"
# placeholders that _normalize_md below matches on.
#
# One of their alternatives is deliberately not ported: a line merely STARTING
# with 自动生成/AI 生成 was dropped whole. In a proposal about code generation
# that eats real sentences ("自动生成的算子在…"), so only self-referential
# phrasings about the document itself are matched here.
#
# This runs on the copy handed to pandoc, never on proposal_final.md. The panel
# reads the .md, and evaluation.decide_verdict caps the verdict at
# revise_resubmit on any leftover placeholder — cleaning the source would hide
# the defect from the judge rather than fix it.

_NOTICE_PAREN = re.compile(
    r"[（(][^（）()]{0,120}"
    r"(?:自动生成|AI\s*生成|人工智能生成|提交前.{0,80}(?:核验|审核|校验|确认))"
    r"[^（）()]{0,80}[）)]",
    re.I,
)

# Self-referential disclaimers — the line names the document AND says it was
# generated. Unambiguous, so trailing text is allowed ("本申请书由人工智能生成，
# 提交前请人工核验。" carries a second clause the original port's anchor missed).
_NOTICE_SELF = re.compile(
    r"^(?:(?:本)?项目申请书|本文档|本文|本申请书)"
    r".*?(?:自动生成|AI\s*生成|人工智能生成).*$",
    re.I,
)

# Looser phrasings, kept strictly anchored: without the end anchor, "提交前…确认"
# would also swallow a genuine 考核方式 sentence such as "提交前完成第三方检测
# 确认，检测报告作为交付物。"
_NOTICE_LINE = re.compile(
    r"^(?:"
    r"(?:注|说明|提示)?[：:]?\s*提交前.*?(?:核验|审核|校验|确认)"
    r"|请.*?(?:人工审核|人工核验|自行核验|核验确认)"
    r")[。；;！!]?$",
    re.I,
)

_PLACEHOLDER_LINE = re.compile(
    r"^(?:"
    r"以下(?:内容)?(?:为)?(?:示例|模板|参考)"
    r"|此处(?:填写|补充|待补充)|待补充|待完善"
    r"|请(?:根据|按)实际情况(?:填写|补充|修改)"
    r")[。；;：:]?$",
    re.I,
)

_EMPTY_PAREN = re.compile(r"[（(]\s*[）)]")


def strip_generator_artifacts(text: str) -> tuple[str, int]:
    """Drop AI disclaimers and placeholder lines from the export copy.

    Returns the cleaned text and how many artefacts were removed, so the caller
    can say so rather than silently shortening a deliverable.
    """
    text, removed = _NOTICE_PAREN.subn("", text)
    kept = []
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped and (_NOTICE_SELF.match(stripped) or _NOTICE_LINE.match(stripped)
                         or _PLACEHOLDER_LINE.match(stripped)):
            removed += 1
            continue
        kept.append(_EMPTY_PAREN.sub("", line))   # "（…自动生成）" can leave "（）"
    return "\n".join(kept), removed


def _normalize_md(text: str, section_names=None) -> str:
    """Polish generator markdown for Word:
    - blank line before pipe tables AND lists (pandoc needs it, else they collapse
      into the preceding paragraph);
    - number the real section headings (一、二、…); demote/drop the writer's repeated
      section title; number ### subheadings （一）（二）… when they lack their own ordinal;
    - break inline enumerations (1)(2)(3)/①②③ onto separate lines;
    - drop the image alt text (else pandoc shows "fig_01") and de-italicise captions.
    `section_names` (from the task) identifies which ## headings are true sections."""
    text = re.sub(r"!\[[^\]]*\]\(", "![](", text)   # strip alt text -> no "fig_01" caption
    expected = list(section_names) if section_names else None
    exp_i = 0
    h2 = 0
    last_section = None       # name of the section just opened (to drop repeats)
    sub = [0]                 # ### counter within the current section
    subsub = [0]              # #### counter within the current ### subsection
    out: list[str] = []
    for ln in text.split("\n"):
        # Two trailing spaces are a Markdown hard line break. The writer ends
        # enumerated items with them out of habit; pandoc then puts a <w:br/> at
        # the end of the list item, which Word draws as an extra empty line
        # inside the bullet. Nothing in a proposal needs a manual line break.
        ln = ln.rstrip()
        prev = out[-1] if out else ""

        # figure caption "*图N：…*" -> plain text (styling done post-conversion)
        cap = re.match(r"^\s*\*\s*(图\s*\d+[：:].*?)\s*\*\s*$", ln)
        if cap:
            out.append(cap.group(1))
            continue
        # un-generated figure placeholder "**[图N：…]**" (>max_figures) -> plain caption
        # (else pandoc leaves a stray ** and the [] brackets)
        ph = re.match(r"^\s*\*\*\s*\[\s*(图\s*\d+[：:].*?)\s*\]\s*\*\*\s*$", ln)
        if ph:
            out.append(ph.group(1))
            continue
        # tables and lists both need a blank line before them (start of the block)
        starts_block = _is_row(ln) or _is_list_item(ln)
        prev_in_block = _is_row(prev) or _is_list_item(prev)
        if starts_block and prev.strip() and not prev_in_block:
            out.append("")
        elif ln.strip() and not _is_row(ln) and _is_row(prev):
            out.append("")   # close a table before following prose

        # level-2 headings (## X) — true section wrapper or a demoted writer heading
        if re.match(r"^##\s+", ln) and not ln.startswith("###"):
            title = ln[2:].strip()
            if title == last_section:
                continue                               # writer repeated the section title -> drop
            if expected is not None:
                if exp_i < len(expected) and title == expected[exp_i]:
                    num = _CN_NUM[exp_i] if exp_i < len(_CN_NUM) else str(exp_i + 1)
                    out.append(f"## {num}、{title}")   # true section wrapper -> numbered
                    exp_i += 1
                    last_section = title
                    sub = [0]; subsub = [0]
                else:
                    out.append(_number_sub(title, sub))  # writer's own heading -> ### + number
                    subsub = [0]
            elif not re.match(r"^[一二三四五六七八九十]+、", title):
                h2 += 1
                num = _CN_NUM[h2 - 1] if h2 - 1 < len(_CN_NUM) else str(h2)
                out.append(f"## {num}、{title}")        # no task list: best-effort number
                last_section = title
                sub = [0]; subsub = [0]
            else:
                out.append(ln)
                last_section = re.sub(r"^[一二三四五六七八九十]+、", "", title)
                sub = [0]; subsub = [0]
            continue

        # genuine ### subheadings from the writer -> number likewise
        if ln.startswith("### "):
            title = ln[4:].strip()
            if title == last_section:
                continue                               # drop repeated section title
            out.append(_number_sub(title, sub))
            subsub = [0]                               # new ### resets the #### counter
            continue

        # #### sub-subheadings -> 1. 2. 3. within the ### parent (一、/（一）/1. 层级)
        if ln.startswith("#### "):
            out.append(_number_subsub(ln[5:].strip(), subsub))
            continue

        # break enumerations on ordinary prose lines (formulas have no bare (digit))
        if ln.strip() and not _is_row(ln) and not ln.startswith("#"):
            ln = _break_enums(ln)
        out.append(ln)
    return "\n".join(out)


def _load_section_names(md_path: Path):
    """Read the sibling task.json (written by the pipeline) to learn the true
    section names, so headings can be numbered without touching the pipeline."""
    tj = md_path.parent / "task.json"
    if tj.exists():
        try:
            d = json.loads(tj.read_text(encoding="utf-8"))
            secs = (d.get("structure") or {}).get("core_sections", [])
            names = [s.get("name", "") for s in secs if s.get("name")]
            return names or None
        except Exception:
            pass
    return None


# schema order of children inside <w:lvl> (for inserting w:suff in the right place)
_LVL_ORDER = ["start", "numFmt", "lvlRestart", "pStyle", "isLgl", "suff",
              "lvlText", "lvlPicBulletId", "legacy", "lvlJc", "pPr", "rPr"]


def _lvl_child(lvl, tag: str):
    el = lvl.find(qn("w:" + tag))
    if el is not None:
        return el
    el = OxmlElement("w:" + tag)
    idx = _LVL_ORDER.index(tag)
    for child in list(lvl):
        cname = child.tag.split("}")[-1]
        if cname in _LVL_ORDER and _LVL_ORDER.index(cname) > idx:
            child.addprevious(el)
            return el
    lvl.append(el)
    return el


# marker style per depth: (numFmt, lvlText-template). {n} is filled with %<level>.
#   0: 1. 2. 3.   1: （1）（2）   2: 1）2）   3: a. b. c.   4: i. ii. iii.
_LVL_STYLE = [
    ("decimal",     "{n}."),
    ("decimal",     "（{n}）"),
    ("decimal",     "{n}）"),
    ("lowerLetter", "{n}."),
    ("lowerRoman",  "{n}."),
]


def _polish_numbering(docx_path: Path):
    """After pandoc, restyle list numbering to a clean multi-level scheme (never a
    bullet glyph ●/○/▪, which is what pandoc emits and we replace):
    - each depth gets its own marker style (1. / （1） / 1） / a. / i.);
    - marker followed by a single space (not a long tab gap);
    - first line indented like body (2 chars), deeper levels stepped in by 2 chars;
    - every list restarts at 1 (Word otherwise continues numbering across all lists
      that share one abstract definition, giving （1）…（40) across the whole doc)."""
    doc = Document(str(docx_path))
    try:
        num_el = doc.part.numbering_part.element
    except Exception:
        return
    for lvl in num_el.findall(".//" + qn("w:lvl")):
        try:
            ilvl = int(lvl.get(qn("w:ilvl")) or 0)
        except ValueError:
            ilvl = 0
        fmt, tmpl = _LVL_STYLE[ilvl % len(_LVL_STYLE)]
        _lvl_child(lvl, "start").set(qn("w:val"), "1")
        _lvl_child(lvl, "numFmt").set(qn("w:val"), fmt)
        _lvl_child(lvl, "lvlText").set(qn("w:val"), tmpl.format(n=f"%{ilvl + 1}"))
        _lvl_child(lvl, "suff").set(qn("w:val"), "space")   # marker + one space

        # drop the bullet symbol font (Symbol/Wingdings) so digits/letters render normally
        rpr = lvl.find(qn("w:rPr"))
        if rpr is not None:
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is not None:
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:cs"), "Times New Roman")
                rfonts.set(qn("w:eastAsia"), "宋体")

        # indent: 2 chars per depth; first line carries the marker
        pPr = _lvl_child(lvl, "pPr")
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        for a in ("w:left", "w:leftChars", "w:hanging", "w:hangingChars",
                  "w:firstLine", "w:firstLineChars", "w:start", "w:startChars"):
            if ind.get(qn(a)) is not None:
                del ind.attrib[qn(a)]
        ind.set(qn("w:leftChars"), str(ilvl * 200))
        ind.set(qn("w:left"), str(ilvl * 480))
        ind.set(qn("w:firstLineChars"), "200")   # 2 chars, same as body first-line indent
        ind.set(qn("w:firstLine"), "480")

    # force EVERY numId instance to restart at 1 (defeats Word's cross-list continue)
    for num in num_el.findall(qn("w:num")):
        for ov in num.findall(qn("w:lvlOverride")):
            num.remove(ov)
        for k in range(9):
            ov = OxmlElement("w:lvlOverride")
            ov.set(qn("w:ilvl"), str(k))
            so = OxmlElement("w:startOverride")
            so.set(qn("w:val"), "1")
            ov.append(so)
            num.append(ov)
    doc.save(str(docx_path))


_CAP_RE = re.compile(r"^图\s*\d+")


def _clear_indent(p):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is not None:
        pPr.remove(ind)


def _polish_figures(docx_path: Path):
    """Center images; render captions centered, upright (no italic) and smaller,
    directly under the image."""
    doc = Document(str(docx_path))
    for p in doc.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):        # a paragraph holding an image
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _clear_indent(p)
            continue
        if _CAP_RE.match(p.text.strip()):                # figure caption
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _clear_indent(p)
            for r in p.runs:
                r.italic = False
                r.bold = False
                r.font.size = Pt(10.5)                   # 五号，比正文(小四/12)小
                rpr = r._r.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = OxmlElement("w:rFonts")
                    rpr.append(rfonts)
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:eastAsia"), "宋体")
    doc.save(str(docx_path))


def _polish_tables(docx_path: Path):
    """Table cells should not carry the body's first-line indent."""
    doc = Document(str(docx_path))
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    ind = pPr.find(qn("w:ind"))
                    if ind is None:
                        ind = OxmlElement("w:ind")
                        pPr.append(ind)
                    ind.set(qn("w:firstLineChars"), "0")
                    ind.set(qn("w:firstLine"), "0")
    doc.save(str(docx_path))


def convert(md_path: Path, out_path: Path | None = None, ref_path: Path | None = None) -> Path:
    md_path = md_path.resolve()
    if out_path is None:
        out_path = md_path.with_suffix(".docx")
    if ref_path is None:
        ref_path = md_path.parents[2] / "assets" / "reference.docx" \
            if len(md_path.parents) >= 3 else md_path.parent / "reference.docx"
    if not ref_path.exists():
        print(f"[ref] building Chinese reference template -> {ref_path}")
        build_reference_docx(ref_path)

    # normalize into a temp file in the same folder so relative image paths still resolve
    section_names = _load_section_names(md_path)
    cleaned, removed = strip_generator_artifacts(md_path.read_text(encoding="utf-8"))
    if removed:
        print(f"[clean] 移除 {removed} 处生成痕迹（免责声明 / 占位行）"
              f"—— 只影响导出的 .docx，{md_path.name} 未改动")
    tmp = md_path.with_name(".__pandoc_tmp.md")
    tmp.write_text(_normalize_md(cleaned, section_names), encoding="utf-8")
    extra = [
        f"--reference-doc={ref_path}",
        f"--resource-path={md_path.parent}",  # resolve figures/xxx.png
    ]
    try:
        pypandoc.convert_file(
            str(tmp), "docx", outputfile=str(out_path),
            format="markdown-implicit_figures+tex_math_single_backslash+tex_math_dollars",
            extra_args=extra,
        )
    finally:
        tmp.unlink(missing_ok=True)
    _polish_numbering(out_path)   # list number spacing + indent
    _polish_figures(out_path)     # centered image + small upright caption
    _polish_tables(out_path)      # table cells: no first-line indent
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md", help="path to proposal_final.md")
    ap.add_argument("--out", default="", help="output .docx (default: alongside md)")
    ap.add_argument("--ref", default="", help="reference.docx template (auto-built if missing)")
    args = ap.parse_args()

    md = Path(args.md)
    if not md.exists():
        sys.exit(f"ERROR: not found: {md}")
    out = convert(md, Path(args.out) if args.out else None, Path(args.ref) if args.ref else None)
    print(f"DONE -> {out}  ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
