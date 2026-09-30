"""Collect every run and judge record into the results workbook (PLAN §7).

    uv run python experiments/collect.py                 # -> runs/exp1/results.xlsx

Machine-produced sheets (tasks, runs, judge_scores, summary) are rebuilt from
the JSON records on every call. The sheets people fill in (human_scores,
blind_map, figure_human) are carried over from the existing workbook if there
is one, so re-collecting never wipes manual entries. agreement is left for the
analysis step once human scores exist.
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict, List

from common import CONDITIONS, EXP, RUNS, load_task, read_json, task_ids

DIMS = [("scientific_quality", "科学质量"), ("innovation", "创新性"), ("feasibility", "可行性"),
        ("alignment", "一致性"), ("impact", "学术影响"), ("compliance", "规范性"), ("clarity", "清晰度")]
VERDICTS = ["recommend_submit", "revise_resubmit", "reject"]

TASK_COLS = ["task_id", "学科", "语言", "资助方 / 项目", "结构来源", "章节数", "字数上限合计",
             "来源链接", "是否新增", "核对人", "核对日期"]
RUN_COLS = ["run_id", "task_id", "条件", "底座 model id", "日期", "commit", "dirty", "状态", "作废原因",
            "字符数", "超限章数", "LLM 调用次数", "输入 token", "输出 token", "耗时(s)",
            "C4 finish_reason", "C1 框架及版本", "C1 轮数", "C1 续跑次数", "规范化删除字数", "缺失章节"]
JUDGE_COLS = (["run_id", "task_id", "条件", "评委", "评委模型", "重复号", "状态"]
              + [n for _, n in DIMS] + ["总分", "判定", "未覆盖要求数", "缺章数", "占位符数", "失败重试次数"])
HUMAN_COLS = ["doc_code", "评审人"] + [n for _, n in DIMS] + ["总分（公式算）", "判定", "用时（分钟）", "低分理由"]
BLIND_COLS = ["doc_code", "run_id", "task_id", "条件"]
FIG_COLS = ["figure_id", "run_id", "评审人", "匹配度档", "表现档", "乱码/不可读", "编造数值", "与正文矛盾",
            "配图评委：匹配度", "配图评委：表现", "配图评委：三类严重问题"]

SECONDARY_STRUCTURE = {"task_004": "二手来源", "task_005": "二手来源", "task_012": "用户提供（非官方模板）"}


# ─────────────────────────────── rows ───────────────────────────────

def task_rows() -> List[List[Any]]:
    rows = []
    for tid in task_ids():
        t = load_task(tid)
        secs = (t.get("structure") or {}).get("core_sections") or []
        limits = [s.get("word_limit") for s in secs if s.get("word_limit")]
        rows.append([tid, t.get("domain", ""), t.get("language", ""), t.get("program", ""),
                     SECONDARY_STRUCTURE.get(tid, "官方模板"), len(secs),
                     sum(limits) if limits else None,
                     "\n".join((t.get("provenance") or {}).get("basis_urls") or []),
                     "是" if tid >= "task_007" else "否", "", ""])
    return rows


def run_rows() -> List[List[Any]]:
    rows = []
    for p in sorted((RUNS / "gen").glob("*/*/r*/run.json")):
        r = read_json(p, {}) or {}
        norm = r.get("normalization") or {}
        rows.append([r.get("run_id"), r.get("task_id"), r.get("condition"), r.get("model"),
                     (r.get("started_at") or "")[:10], (r.get("commit") or "")[:10], r.get("dirty"),
                     r.get("status"), "; ".join(r.get("void_reasons") or []), r.get("char_count"),
                     len(r.get("over_limit") or []), r.get("llm_calls"), r.get("prompt_tokens"),
                     r.get("completion_tokens"), r.get("elapsed_seconds"), r.get("finish_reason"),
                     r.get("framework"), r.get("turns"), r.get("continuations"),
                     norm.get("removed_chars"), "; ".join(norm.get("missing_sections") or [])])
    return rows


def judge_rows() -> List[List[Any]]:
    rows = []
    for p in sorted((RUNS / "judge").glob("*/*/*/r*_j*.json")):
        j = read_json(p, {}) or {}
        res = j.get("result") or {}
        scores = res.get("scores") or {}
        rows.append([j.get("run_id"), j.get("task_id"), j.get("condition"), j.get("judge"),
                     j.get("judge_model"), j.get("judge_rep"), j.get("status")]
                    + [scores.get(k) for k, _ in DIMS]
                    + [res.get("overall_100"), res.get("verdict"),
                       sum(1 for c in res.get("requirement_coverage") or [] if not c.get("covered", True)),
                       sum(1 for s in res.get("section_checklist") or [] if not s.get("present", True)),
                       len(res.get("placeholders") or []), len(j.get("failed_attempts") or [])])
    return rows


def judge_names() -> List[str]:
    return [j["name"] for j in read_json(EXP / "judges.json", {})["judges"]]


# ─────────────────────────────── workbook ───────────────────────────────

def build(out: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    keep: Dict[str, List[List[Any]]] = {}
    if out.exists():                      # manual sheets survive a re-collect
        old = load_workbook(out)
        for name in ("human_scores", "blind_map", "figure_human"):
            if name in old.sheetnames:
                keep[name] = [list(r) for r in old[name].iter_rows(min_row=2, values_only=True)
                              if any(v is not None for v in r)]

    wb = Workbook()
    wb.remove(wb.active)
    head_font, body_font = Font(name="Arial", bold=True), Font(name="Arial")
    input_fill = PatternFill("solid", start_color="FFFF00")

    def sheet(name: str, cols: List[str], rows: List[List[Any]], input_sheet: bool = False):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = head_font
            if input_sheet:
                c.fill = input_fill
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(col)) * 2 + 2))
        ws.freeze_panes = "A2"
        return ws

    readme = wb.create_sheet("说明")
    for line in [
        ["结果表：实验 1（架构对比）。方案见 experiments/PLAN.md §7。"],
        ["tasks、runs、judge_scores、summary 由 experiments/collect.py 从 runs/exp1 下的记录自动生成，重跑即覆盖，不要手改。"],
        ["human_scores、blind_map、figure_human 表头为黄色，是人工填写的表；重跑 collect.py 时会原样保留。"],
        ["human_scores 填法：每篇 × 每名评审人一行；7 维各打 1–10 整数；总分列自动按权重计算；判定从下拉选择。"],
        ["human_scores 示例行（不要填进 human_scores）：D07 | 评审人A | 8 | 7 | 7 | 9 | 8 | 7 | 8 | （自动） | recommend_submit | 55 | "],
        ["总分权重与评审代码 evaluation.WEIGHTS 相同：科学质量 0.20、创新性 0.18、可行性 0.16、一致性 0.14、学术影响 0.12、规范性 0.11、清晰度 0.09。"],
    ]:
        readme.append(line)
    readme.column_dimensions["A"].width = 120
    for row in readme.iter_rows():
        for c in row:
            c.font = body_font
            c.alignment = Alignment(wrap_text=True)

    sheet("tasks", TASK_COLS, task_rows())
    sheet("runs", RUN_COLS, run_rows())
    js = sheet("judge_scores", JUDGE_COLS, judge_rows())

    # summary: task x condition, every figure a formula over judge_scores / runs
    judges = judge_names()
    n_js = max(js.max_row, 2)
    col = {name: get_column_letter(JUDGE_COLS.index(name) + 1) for name in ("task_id", "条件", "评委", "总分", "状态")}
    rng = lambda c: f"judge_scores!${col[c]}$2:${col[c]}${n_js}"          # noqa: E731
    n_runs = max(len(run_rows()) + 1, 2)
    rcol = {name: get_column_letter(RUN_COLS.index(name) + 1) for name in ("task_id", "条件", "字符数", "状态")}
    rrng = lambda c: f"runs!${rcol[c]}$2:${rcol[c]}${n_runs}"            # noqa: E731
    sm_cols = ["task_id", "条件"] + [f"{j} 总分" for j in judges] + ["评审团均分", "人工均分", "字符数"]
    ss = sheet("summary", sm_cols, [])
    r = 2
    for tid in task_ids():
        for cond in CONDITIONS:
            ss.cell(r, 1, tid).font = body_font
            ss.cell(r, 2, cond).font = body_font
            for k, jn in enumerate(judges):
                ss.cell(r, 3 + k, (f'=IFERROR(AVERAGEIFS({rng("总分")},{rng("task_id")},$A{r},'
                                   f'{rng("条件")},$B{r},{rng("评委")},"{jn}",{rng("状态")},"valid"),"")'))
            first, last = get_column_letter(3), get_column_letter(2 + len(judges))
            ss.cell(r, 3 + len(judges), f'=IFERROR(AVERAGE({first}{r}:{last}{r}),"")')
            ss.cell(r, 4 + len(judges), "")      # filled by the analysis step via blind_map
            ss.cell(r, 5 + len(judges), (f'=IFERROR(AVERAGEIFS({rrng("字符数")},{rrng("task_id")},$A{r},'
                                         f'{rrng("条件")},$B{r},{rrng("状态")},"valid"),"")'))
            for c in ss[r]:
                c.font = body_font
            r += 1

    hs = sheet("human_scores", HUMAN_COLS, keep.get("human_scores", []), input_sheet=True)
    weights = [0.20, 0.18, 0.16, 0.14, 0.12, 0.11, 0.09]
    first_dim = 3
    total_col = first_dim + len(DIMS)
    for row in range(2, max(hs.max_row, 101) + 1):
        terms = "+".join(f"{get_column_letter(first_dim + i)}{row}*{w}" for i, w in enumerate(weights))
        dims = f"{get_column_letter(first_dim)}{row}:{get_column_letter(first_dim + len(DIMS) - 1)}{row}"
        hs.cell(row, total_col, f'=IF(COUNT({dims})=7,ROUND(({terms})*10,1),"")').font = body_font
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True)
    verdict_dv = DataValidation(type="list", formula1=f'"{",".join(VERDICTS)}"', allow_blank=True)
    hs.add_data_validation(score_dv)
    hs.add_data_validation(verdict_dv)
    score_dv.add(f"{get_column_letter(first_dim)}2:{get_column_letter(first_dim + len(DIMS) - 1)}200")
    verdict_dv.add(f"{get_column_letter(total_col + 1)}2:{get_column_letter(total_col + 1)}200")

    sheet("blind_map", BLIND_COLS, keep.get("blind_map", []), input_sheet=True)
    agreement = sheet("agreement", ["指标", "总分"] + [n for _, n in DIMS], [])
    for name in ["人–人 ICC(2,1)", "评审团–人 Spearman ρ"] + [f"{j}–人 Spearman ρ" for j in judges] + \
                ["同一 task 内条件排序 Kendall τ", "判定一致性（加权 κ）", "系统性偏差（评审团 − 人）"]:
        agreement.append([name])
    fh = sheet("figure_human", FIG_COLS, keep.get("figure_human", []), input_sheet=True)
    fdv = DataValidation(type="list", formula1='"good,partial,mismatch"', allow_blank=True)
    qdv = DataValidation(type="list", formula1='"good,acceptable,poor"', allow_blank=True)
    bdv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    for dv in (fdv, qdv, bdv):
        fh.add_data_validation(dv)
    fdv.add("D2:D500")
    qdv.add("E2:E500")
    bdv.add("F2:H500")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"-> {out}  (runs {len(run_rows())}, judge records {len(judge_rows())})")
    print("   Formulas have no cached values until the file is opened in Excel/LibreOffice or recalculated.")


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else RUNS / "results.xlsx")
